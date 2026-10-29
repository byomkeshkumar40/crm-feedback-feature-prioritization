"""Automated tests for the SYNTHETIC CRM feedback pipeline. Run: pytest -q  (after run_pipeline.sh)."""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import analyze_feedback as af  # noqa: E402
import calculate_rice as cr  # noqa: E402
import clean_data  # noqa: E402
import generate_data as gd  # noqa: E402
import themes  # noqa: E402
import validate_data as vd  # noqa: E402

PROC = ROOT / "data" / "processed"


@pytest.fixture(scope="module")
def clean():
    return pd.read_csv(PROC / "feedback_clean.csv", parse_dates=["submission_date"])


@pytest.fixture(scope="module")
def cust():
    return pd.read_csv(ROOT / "data" / "raw" / "customers.csv")


def test_exactly_1500_records(clean):
    assert len(clean) == 1500


def test_generator_deterministic():
    a, _ = gd.generate()
    b, _ = gd.generate()
    pd.testing.assert_frame_equal(a, b)


def test_no_duplicate_ids(clean):
    assert clean.feedback_id.is_unique


def test_customer_relationships(clean, cust):
    assert clean.customer_id.isin(cust.customer_id).all()
    m = clean.merge(cust, on="customer_id", suffixes=("", "_c"))
    assert (m.subscription_plan == m.subscription_plan_c).all()
    assert (m.customer_segment == m.customer_segment_c).all()
    assert (m.submission_date >= pd.to_datetime(m.signup_date)).all()


def test_categorical_validity(clean):
    for col, allowed in clean_data.CANON.items():
        extra = set(clean[col].dropna()) - set(allowed) - {"Unknown"}
        assert not extra, (col, extra)


def test_area_category_consistency(clean):
    g = clean.groupby("issue_theme")[["feature_area", "feedback_category"]].nunique()
    assert (g == 1).all().all()
    assert (clean[clean.feedback_category == "Feature Request"].feedback_source.value_counts(normalize=True)["Feature Request"] > 0.3)


def test_text_variety(clean):
    assert clean.feedback_text.nunique() / len(clean) >= 0.95


def test_theme_keys_unambiguous():
    df, _ = gd.generate(with_truth=True)
    tags = df.feedback_text.map(themes.tag_text)
    assert (tags.map(len) == 1).all()
    assert (tags.str[0] == df["_theme"]).all()


def test_cleaning_fixes_noise():
    raw = pd.read_csv(ROOT / "data" / "raw" / "feedback_raw.csv", dtype=str)
    out, log = clean_data.clean(raw)
    assert len(out) == len(raw) == 1500
    assert out.submission_date.notna().all()
    assert log["non_iso_dates_converted"] > 0 and log["missing_customer_impact_before"] > 0
    assert (out.customer_impact == "Unknown").sum() == log["customer_impact_imputed_as_Unknown"]


def test_validator_flags_bad_data(cust, clean):
    bad = clean.copy()
    bad.loc[0, "feedback_id"] = bad.loc[1, "feedback_id"]
    bad.loc[2, "customer_id"] = "CUST-9999"
    bad.loc[3, "sentiment"] = "Angry"
    failed = {n for n, ok, _ in vd.run_checks(bad, cust, "clean") if not ok}
    assert {"no duplicate feedback_id", "all customer_ids exist in customer master", "sentiment: only valid values"} <= failed
    assert all(ok for _, ok, _ in vd.run_checks(clean, cust, "clean"))


def test_counts_match_pandas(clean):
    _, res, tabs = af.analyse()
    assert res["total_records"] == len(clean)
    assert tabs["by_feature_area"].feedback_records.sum() == 1500
    assert tabs["by_feature_area"].feedback_records.to_dict() == clean.feature_area.value_counts().to_dict()
    assert tabs["sentiment"].feedback_records.sum() == 1500
    assert tabs["by_plan"].accounts_in_base.sum() == 420
    assert tabs["high_sev_by_area"].high_critical_records.sum() == clean.severity.isin(["High", "Critical"]).sum()
    assert res["unique_customers"] == clean.customer_id.nunique() <= res["total_records"]


def test_rice_math_and_ranks():
    base, ranks, scores, stab, meta = cr.run()
    assert len(base) == 12 and base.feature_id.is_unique
    exp = (base.reach * base.impact * base.confidence / base.effort).round(2)
    assert np.allclose(base.rice_score, exp, atol=0.011)
    assert (base.sort_values("rice_score", ascending=False).priority_rank.values == sorted(base.priority_rank)).all()
    assert base.priority_rank.is_unique  # no ties in base case
    assert (base.reach <= meta["accounts"]).all()
    assert set(base.impact) <= {0.25, 0.5, 1, 2, 3}
    assert (ranks["Base case"] == base.set_index("feature_id").priority_rank.reindex(ranks.index)).all()
    assert abs(stab.prob_top3_pct.sum() - 300) < 0.5  # 3 features in the top 3 per trial; values rounded to 0.1


def test_rice_rounding_half_up():
    assert cr.impact_from_signal(0.55) == 3 and cr.impact_from_signal(0.1) == 0.25
    assert cr.confidence_rule(150, 4, 3) == 1.0 and cr.confidence_rule(10, 1, 1) == 0.5


def test_rice_files_consistent_with_assumptions(clean):
    csv = pd.read_csv(PROC / "rice" / "rice_scores.csv")
    base, *_ = cr.run()
    assert csv.set_index("feature_id").rice_score.round(2).equals(base.set_index("feature_id").rice_score.round(2).reindex(csv.feature_id))


def test_excel_workbook():
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "excel" / "CRM_Feedback_Analysis.xlsx")
    assert wb.sheetnames == ["Executive Summary", "Raw Feedback", "Cleaned Feedback", "Feedback Analysis",
                             "Feature Prioritization", "RICE Calculations", "Product Recommendations"]
    assert wb["Cleaned Feedback"].max_row == 1501
    assert wb["Cleaned Feedback"].freeze_panes == "A2" and wb["Cleaned Feedback"].auto_filter.ref


def test_reports_consistent_with_results():
    base = pd.read_csv(PROC / "rice" / "rice_scores.csv")
    top = base.sort_values("priority_rank").iloc[0]
    prd = (ROOT / "reports" / "product_requirements_document.md").read_text()
    assert top.feature_name in prd and f"{int(top.supporting_feedback_count)} feedback records" in prd
    ins = (ROOT / "reports" / "product_insights.md").read_text()
    assert f"{top.rice_score:g}" in ins and "SYNTHETIC" in ins
    for f in ["feedback_overview.png", "customer_pain_points.png", "feature_prioritization.png"]:
        assert (ROOT / "dashboards" / "screenshots" / f).stat().st_size > 30_000
