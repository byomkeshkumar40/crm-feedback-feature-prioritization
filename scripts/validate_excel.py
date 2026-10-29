"""Validate the Excel workbook: structure with OpenPyXL, formula results by independent recalculation (formulas lib)."""
import sys
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
ROOT = Path(__file__).resolve().parents[1]
XL = ROOT / "excel" / "CRM_Feedback_Analysis.xlsx"
SHEETS = ["Executive Summary", "Raw Feedback", "Cleaned Feedback", "Feedback Analysis", "Feature Prioritization",
          "RICE Calculations", "Product Recommendations"]


def structural():
    wb = load_workbook(XL)
    out = [("sheet names/order", wb.sheetnames == SHEETS, wb.sheetnames)]
    for n in ("Raw Feedback", "Cleaned Feedback"):
        out.append((f"{n}: 1500 data rows", wb[n].max_row == 1501, wb[n].max_row))
    for n in ("Raw Feedback", "Cleaned Feedback", "Feature Prioritization", "RICE Calculations", "Product Recommendations"):
        out.append((f"{n}: filter + frozen header", bool(wb[n].auto_filter.ref) and bool(wb[n].freeze_panes), (wb[n].auto_filter.ref, wb[n].freeze_panes)))
    for n in ("Feature Prioritization", "RICE Calculations"):
        out.append((f"{n}: conditional formatting", len(wb[n].conditional_formatting) > 0, len(wb[n].conditional_formatting)))
    nform = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    out.append(("contains working formulas", nform > 100, nform))
    return out


def recalculated():
    import formulas
    warnings.filterwarnings("ignore")
    xl = formulas.ExcelModel().loads(str(XL)).finish()
    sol = xl.calculate()
    def val(sheet, ref):
        key = f"'[{XL.name}]{sheet.upper()}'!{ref}"
        v = sol[key].value[0, 0]
        return v
    return val


def main():
    res = structural()
    val = recalculated()
    df = pd.read_csv(ROOT / "data" / "processed" / "feedback_clean.csv")
    rice = pd.read_csv(ROOT / "data" / "processed" / "rice" / "rice_scores.csv").set_index("feature_id")
    ES = "Executive Summary"
    chk = lambda n, a, b, tol=1e-6: res.append((n, abs(float(a) - float(b)) <= tol, f"excel={a} python={b}"))  # noqa: E731
    chk("KPI total feedback", val(ES, "B6"), len(df))
    chk("KPI negative %", val(ES, "B8"), (df.sentiment == "Negative").mean())
    chk("KPI high/critical %", val(ES, "B9"), df.severity.isin(["High", "Critical"]).mean())
    chk("KPI feature requests", val(ES, "B10"), (df.feedback_category == "Feature Request").sum())
    for f in rice.index:
        r = 9 + sorted(rice.index).index(f)
        chk(f"RICE score {f}", val("RICE Calculations", f"I{r}"), rice.loc[f, "rice_score"], 0.011)
        chk(f"RICE rank {f}", val("RICE Calculations", f"J{r}"), rice.loc[f, "priority_rank"])
        chk(f"Reach {f}", val("RICE Calculations", f"E{r}"), rice.loc[f, "reach"])
    top = rice.sort_values("priority_rank").iloc[0]
    res.append(("Exec summary top-1 name matches", str(val(ES, "A14")) == top.feature_name, str(val(ES, "A14"))))
    area = df.feature_area.value_counts()
    for i, (k, v) in enumerate(area.items()):
        pass
    wb = load_workbook(XL)
    wa = wb["Feedback Analysis"]
    r = 6
    for i in range(r, r + 6):
        k = wa.cell(row=i, column=1).value
        chk(f"COUNTIF feature_area {k}", val("Feedback Analysis", f"B{i}"), area[k])
    for n, ok, d in res:
        print(("PASS" if ok else "FAIL"), n, d)
    print(f"\n{sum(o for _, o, _ in res)}/{len(res)} workbook checks passed")
    sys.exit(0 if all(o for _, o, _ in res) else 1)


if __name__ == "__main__":
    main()
