"""Build excel/CRM_Feedback_Analysis.xlsx (SYNTHETIC data) with OpenPyXL."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analyze_feedback import analyse  # noqa: E402
from calculate_rice import run as run_rice, BASE, WINDOW  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "CRM_Feedback_Analysis.xlsx"
NAVY, BLUE, LIGHT = "1F3A5F", "2F6FDE", "EEF3FB"
HFONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HFILL = PatternFill("solid", fgColor=NAVY)
EST_FILL = PatternFill("solid", fgColor="FFF4D6")   # estimates
OBS_FILL = PatternFill("solid", fgColor="E3F2E8")   # observed
thin = Side(style="thin", color="D0D7E2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BANNER = "SYNTHETIC DATA — all records are randomly generated for a fictional CRM company; no real customers."


def title(ws, text, sub=None, width=8):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = BANNER
    ws["A2"].font = Font(italic=True, color="B42318", size=10)
    if sub:
        ws["A3"] = sub
        ws["A3"].font = Font(color="555555", size=10)


def header(ws, row, cols, start=1):
    for j, c in enumerate(cols):
        cell = ws.cell(row=row, column=start + j, value=c)
        cell.font, cell.fill, cell.border = HFONT, HFILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


def box(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = BORDER


def dump_df(ws, df, r0=1, datecols=()):
    header(ws, r0, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif pd.isna(v) if not isinstance(v, (list, tuple)) else False:
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            if df.columns[j - 1] in datecols:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.auto_filter.ref = f"A{r0}:{L(len(df.columns))}{r0 + len(df)}"


def build():
    df, res, tabs = analyse()
    raw = pd.read_csv(ROOT / "data" / "raw" / "feedback_raw.csv", dtype=str)
    base, ranks, scores, stab, meta = run_rice()
    n = len(df)
    wb = Workbook()
    # ---------- Raw / Cleaned ----------
    wr = wb.active; wr.title = "Raw Feedback"
    dump_df(wr, raw)
    widths(wr, [11, 11, 13, 16, 14, 15, 20, 18, 80, 11, 10, 22])
    wc = wb.create_sheet("Cleaned Feedback")
    cl = df.copy()
    cl = cl[["feedback_id", "customer_id", "submission_date", "feedback_source", "subscription_plan", "customer_segment",
             "feature_area", "feedback_category", "feedback_text", "sentiment", "severity", "customer_impact",
             "issue_theme", "issue_theme_label", "month"]]
    dump_df(wc, cl, datecols=("submission_date",))
    widths(wc, [11, 11, 13, 16, 14, 15, 20, 18, 80, 11, 10, 22, 16, 30, 10])
    last = n + 1
    rng = lambda col: f"'Cleaned Feedback'!${col}$2:${col}${last}"  # noqa: E731
    # ---------- Feedback Analysis (COUNTIF formulas) ----------
    wa = wb.create_sheet("Feedback Analysis")
    title(wa, "Feedback Analysis", "Record counts are live COUNTIF formulas on 'Cleaned Feedback'. "
          "Unique-customer counts are computed in Python (static). Volume ≠ customer prevalence.")
    r = 5
    blocks = [("Feature Area", "G", tabs["by_feature_area"]), ("Feedback Category", "H", tabs["by_category"]),
              ("Sentiment", "J", tabs["sentiment"]), ("Severity", "K", tabs["severity"]),
              ("Subscription Plan", "E", tabs["by_plan"]), ("Customer Segment", "F", tabs["by_segment"]),
              ("Feedback Source", "D", tabs["by_source"])]
    for name, col, t in blocks:
        header(wa, r, [name, "Feedback Records", "% of Records", "Unique Customers (static)"])
        for i, (k, row) in enumerate(t.iterrows(), r + 1):
            wa.cell(row=i, column=1, value=k)
            wa.cell(row=i, column=2, value=f'=COUNTIF({rng(col)},A{i})')
            c = wa.cell(row=i, column=3, value=f"=B{i}/COUNTA({rng('A')})"); c.number_format = "0.0%"
            wa.cell(row=i, column=4, value=int(row.unique_customers))
        box(wa, r, r + len(t), 1, 4)
        wa.conditional_formatting.add(f"B{r + 1}:B{r + len(t)}", DataBarRule(start_type="num", start_value=0, end_type="max", color=BLUE))
        r += len(t) + 3
    header(wa, r, ["Month", "Feedback Records", "Unique Customers (static)"])
    mt = tabs["monthly_trend"]
    for i, (k, row) in enumerate(mt.iterrows(), r + 1):
        wa.cell(row=i, column=1, value=k)
        wa.cell(row=i, column=2, value=f'=COUNTIF({rng("O")},A{i})')
        wa.cell(row=i, column=3, value=int(row.unique_customers))
    box(wa, r, r + len(mt), 1, 3)
    r += len(mt) + 3
    header(wa, r, ["Top Complaint Themes (non-request)", "Feature Area", "Category", "Feedback Records", "Unique Customers",
                   "Negative %", "High/Critical %"])
    for i, (k, row) in enumerate(tabs["top_complaints"].iterrows(), r + 1):
        vals = [row.theme, row.feature_area, row.category, f'=COUNTIF({rng("M")},"{k}")', int(row.unique_customers),
                row.negative_pct / 100, row.high_critical_pct / 100]
        for j, v in enumerate(vals, 1):
            c = wa.cell(row=i, column=j, value=v)
            if j >= 6: c.number_format = "0.0%"
    box(wa, r, r + 10, 1, 7)
    r += 13
    header(wa, r, ["Most Requested Features (theme)", "Feature Area", "Category", "Feedback Records", "Unique Customers"])
    for i, (k, row) in enumerate(tabs["top_requests"].iterrows(), r + 1):
        for j, v in enumerate([row.theme, row.feature_area, row.category, f'=COUNTIF({rng("M")},"{k}")',
                               int(row.unique_customers)], 1):
            wa.cell(row=i, column=j, value=v)
    box(wa, r, r + 10, 1, 5)
    r += 13
    header(wa, r, ["High/Critical Issues by Feature Area", "High+Critical Records", "% of Area Records", "Unique Customers (static)"])
    hs = tabs["high_sev_by_area"]
    for i, (k, row) in enumerate(hs.iterrows(), r + 1):
        wa.cell(row=i, column=1, value=k)
        wa.cell(row=i, column=2, value=f'=COUNTIFS({rng("G")},A{i},{rng("K")},"High")+COUNTIFS({rng("G")},A{i},{rng("K")},"Critical")')
        c = wa.cell(row=i, column=3, value=f'=B{i}/COUNTIF({rng("G")},A{i})'); c.number_format = "0.0%"
        wa.cell(row=i, column=4, value=int(row.unique_customers))
    box(wa, r, r + len(hs), 1, 4)
    widths(wa, [38, 20, 20, 22, 18, 12, 14])
    # ---------- RICE Calculations ----------
    wx = wb.create_sheet("RICE Calculations")
    title(wx, "RICE Calculations", f"Reach window: {WINDOW}. RICE = Reach × Impact × Confidence ÷ Effort. "
          "Green = observed from data; yellow = analyst ESTIMATE (editable — scores and ranks recalculate).")
    wx["A5"] = "Reporting multiplier (estimate)"; wx["B5"] = BASE["reporting_multiplier"]; wx["B5"].fill = EST_FILL
    wx["C5"] = "Assumed # of affected accounts represented by each account that reported (feedback under-reports)."
    wx["A6"] = "Account base / reach cap (observed)"; wx["B6"] = meta["accounts"]; wx["B6"].fill = OBS_FILL
    wx["C6"] = "Accounts in synthetic customer master; Reach cannot exceed this."
    cols = ["Feature ID", "Feature Name", "Supporting Feedback (observed)", "Unique Affected Customers (observed)",
            "Reach (est. accounts/12 mo)", "Impact (estimate)", "Confidence (estimate)", "Effort (person-months, estimate)",
            "RICE Score", "Priority Rank", "Python RICE (cross-check)", "How estimates were chosen"]
    h = 8
    header(wx, h, cols)
    b = base.sort_values("feature_id").reset_index(drop=True)
    for i, row in enumerate(b.itertuples(), h + 1):
        vals = [row.feature_id, row.feature_name, row.supporting_feedback_count, row.unique_affected_customers,
                f"=MIN($B$6,ROUND(D{i}*$B$5,0))", row.impact, row.confidence, row.effort,
                f"=ROUND(E{i}*F{i}*G{i}/H{i},2)", f"=RANK(I{i},$I${h + 1}:$I${h + len(b)})", row.rice_score, row.estimate_rationale]
        for j, v in enumerate(vals, 1):
            c = wx.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j in (3, 4): c.fill = OBS_FILL
            if j in (6, 7, 8): c.fill = EST_FILL
            if j == 7: c.number_format = "0.00"
    box(wx, h, h + len(b), 1, 12)
    first, lastr = h + 1, h + len(b)
    wx.conditional_formatting.add(f"I{first}:I{lastr}", ColorScaleRule(start_type="min", start_color="F8D7DA", mid_type="percentile",
                                  mid_value=50, mid_color="FFF3CD", end_type="max", end_color="C6E8CF"))
    wx.freeze_panes = f"C{h + 1}"
    wx.auto_filter.ref = f"A{h}:L{lastr}"
    s = lastr + 3
    wx.cell(row=s, column=1, value="Sensitivity analysis — priority rank under alternative assumptions (computed in Python)").font = Font(bold=True, size=12, color=NAVY)
    rk = ranks.reset_index().drop(columns=["rank_min", "rank_max"])
    rk["Best rank"] = ranks.rank_min.values; rk["Worst rank"] = ranks.rank_max.values
    header(wx, s + 1, list(rk.columns))
    for i, row in enumerate(rk.itertuples(index=False), s + 2):
        for j, v in enumerate(row, 1):
            wx.cell(row=i, column=j, value=v if not hasattr(v, "item") else v.item())
    box(wx, s + 1, s + 1 + len(rk), 1, len(rk.columns))
    s2 = s + len(rk) + 4
    wx.cell(row=s2, column=1, value=f"Monte Carlo rank stability — {meta['mc_trials']} trials, each of R, I, C, E perturbed ±30% (uniform, seed 7)").font = Font(bold=True, size=12, color=NAVY)
    st = stab.reset_index().rename(columns={"feature_id": "Feature ID"})
    header(wx, s2 + 1, list(st.columns))
    for i, row in enumerate(st.itertuples(index=False), s2 + 2):
        for j, v in enumerate(row, 1):
            wx.cell(row=i, column=j, value=v if not hasattr(v, "item") else v.item())
    box(wx, s2 + 1, s2 + 1 + len(st), 1, len(st.columns))
    widths(wx, [30, 40, 16, 16, 16, 12, 14, 16, 12, 12, 16, 110])
    wx.row_dimensions[h].height = 48
    # ---------- Feature Prioritization ----------
    wf = wb.create_sheet("Feature Prioritization")
    title(wf, "Feature Prioritization", "Scores/ranks are formulas linked to 'RICE Calculations'. RICE scores are decision aids built on estimates, not objective facts.")
    cols = ["Feature ID", "Feature Name", "Customer Problem", "Supporting Feedback Count", "Unique Affected Customers",
            "Reach", "Impact", "Confidence", "Effort (pm)", "RICE Score", "Priority Rank", "Recommendation"]
    header(wf, 5, cols)
    for i, row in enumerate(b.itertuples(), 6):
        rr = i + h + 1 - 6 + 0  # same row index on RICE sheet
        rr = h + 1 + (i - 6)
        vals = [row.feature_id, row.feature_name, row.customer_problem, f"='RICE Calculations'!C{rr}", f"='RICE Calculations'!D{rr}",
                f"='RICE Calculations'!E{rr}", f"='RICE Calculations'!F{rr}", f"='RICE Calculations'!G{rr}",
                f"='RICE Calculations'!H{rr}", f"='RICE Calculations'!I{rr}", f"='RICE Calculations'!J{rr}", row.recommendation]
        for j, v in enumerate(vals, 1):
            c = wf.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    box(wf, 5, 5 + len(b), 1, 12)
    wf.conditional_formatting.add(f"J6:J{5 + len(b)}", ColorScaleRule(start_type="min", start_color="F8D7DA", mid_type="percentile",
                                  mid_value=50, mid_color="FFF3CD", end_type="max", end_color="C6E8CF"))
    wf.freeze_panes = "C6"; wf.auto_filter.ref = f"A5:L{5 + len(b)}"
    widths(wf, [11, 38, 55, 14, 14, 10, 9, 11, 10, 11, 10, 32])
    wf.row_dimensions[5].height = 36
    # ---------- Product Recommendations ----------
    wp = wb.create_sheet("Product Recommendations")
    title(wp, "Product Recommendations", "Hypotheses for the product team, derived from RICE ranks and the feedback analysis; not validated outcomes.")
    header(wp, 5, ["Priority", "Feature", "Recommendation", "Why (observed data)", "Open risk / what to validate"])
    top = base.head(12)
    for i, row in enumerate(top.itertuples(), 6):
        why = (f"{row.supporting_feedback_count} records from {row.unique_affected_customers} accounts; "
               f"{row.combined_signal:.0%} High/Critical or blocking; RICE {row.rice_score:g}")
        risk = "Rank is sensitive to effort estimate" if ranks.loc[row.feature_id, "rank_max"] - ranks.loc[row.feature_id, "rank_min"] >= 3 else "Rank stable across scenarios"
        for j, v in enumerate([row.priority_rank, f"{row.feature_id} {row.feature_name}", row.recommendation, why, risk], 1):
            wp.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    box(wp, 5, 5 + len(top), 1, 5)
    wp.freeze_panes = "A6"; wp.auto_filter.ref = f"A5:E{5 + len(top)}"
    widths(wp, [9, 46, 34, 70, 36])
    # ---------- Executive Summary ----------
    we = wb.create_sheet("Executive Summary", 0)
    title(we, "CRM Customer Feedback Analysis & Feature Prioritization", "Executive summary — fictional B2B SaaS CRM, 12-month window (Oct 2025 – Sep 2026)")
    kp = [("Total feedback records", f"=COUNTA({rng('A')})", "#,##0"),
          ("Unique customers (static)", res["unique_customers"], "#,##0"),
          ("Negative feedback %", f'=COUNTIF({rng("J")},"Negative")/COUNTA({rng("A")})', "0.0%"),
          ("High/Critical severity %", f'=(COUNTIF({rng("K")},"High")+COUNTIF({rng("K")},"Critical"))/COUNTA({rng("A")})', "0.0%"),
          ("Feature requests", f'=COUNTIF({rng("H")},"Feature Request")', "#,##0")]
    header(we, 5, ["KPI", "Value"])
    for i, (k, v, f) in enumerate(kp, 6):
        we.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = we.cell(row=i, column=2, value=v); c.number_format = f; c.font = Font(bold=True, size=13, color=BLUE)
    box(we, 5, 10, 1, 2)
    header(we, 13, ["Top 3 RICE priorities", "RICE Score", "Recommendation"])
    rc = "'RICE Calculations'"
    for k in (1, 2, 3):
        i = 13 + k
        we.cell(row=i, column=1, value=f"=INDEX({rc}!$B${first}:$B${lastr},MATCH({k},{rc}!$J${first}:$J${lastr},0))")
        we.cell(row=i, column=2, value=f"=INDEX({rc}!$I${first}:$I${lastr},MATCH({k},{rc}!$J${first}:$J${lastr},0))")
        we.cell(row=i, column=3, value=base.loc[base.priority_rank == k, "recommendation"].iloc[0])
    box(we, 13, 16, 1, 3)
    topc, tops = tabs["top_complaints"].iloc[0], tabs["by_feature_area"]
    notes = ["Observed (synthetic data):",
             f"• {res['total_records']:,} records from {res['unique_customers']} of {res['customer_master_accounts']} accounts ({res['pct_accounts_with_feedback']}%).",
             f"• Most-reported complaint theme: {topc.theme} ({int(topc.feedback_records)} records, {int(topc.unique_customers)} accounts).",
             f"• Largest feature area by volume: {tops.index[0]} ({int(tops.feedback_records.iloc[0])} records).",
             "Interpretation: volume reflects reports, not prevalence; the 420-account base is the only denominator.",
             "Hypothesis: fixing top-ranked reliability issues (import, workflow monitoring) may reduce support load — to be validated.",
             "Limitations: synthetic data, estimated Impact/Confidence/Effort, RICE ranks are not objective facts."]
    for i, t in enumerate(notes, 19):
        we.cell(row=i, column=1, value=t).font = Font(bold=(i == 19))
    widths(we, [48, 22, 34])
    order = ["Executive Summary", "Raw Feedback", "Cleaned Feedback", "Feedback Analysis", "Feature Prioritization", "RICE Calculations", "Product Recommendations"]
    wb._sheets = [wb[s] for s in order]
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False if ws.title in ("Executive Summary",) else True
    return wb


def main():
    OUT.parent.mkdir(exist_ok=True)
    build().save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
